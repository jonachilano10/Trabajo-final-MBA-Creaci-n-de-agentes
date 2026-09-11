from __future__ import annotations

import json
import re
import unicodedata
import argparse
from collections import Counter, defaultdict
from datetime import date, datetime, time
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "2-Trabajo a entregar" / "Archivos excel para cargar"
FILES = {
    "avisos": INPUT / "33-Avisos MG (1).XLSX",
    "notificaciones": INPUT / "33-Notificaciones (1).XLSX",
    "mantenimiento": INPUT / "33-TAREAS EN SALA SEMANA 33 (1).xlsx",
}
START = date(2026, 8, 10)
END = date(2026, 8, 16)
WEEK = 33
YEAR = 2026
TAGS = [
    "P-MG11", "P-MG12", "P-MG13",
    "P-MG21", "P-MG22", "P-MG23",
    "P-MG31", "P-MG32", "P-MG33",
    "P-MG41", "P-MG42", "P-MG43",
]
TAG_SET = set(TAGS)
ROOM = {tag: int(tag[4]) for tag in TAGS}
SERIES = {tag: ("A" if ROOM[tag] <= 2 else "B") for tag in TAGS}


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def normalize(value):
    text = clean(value).lower()
    text = unicodedata.normalize("NFKD", text)
    return "".join(c for c in text if not unicodedata.combining(c))


def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = clean(value)
    if not s:
        return None
    if " " in s and s[:4].isdigit():
        s = s.split()[0]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def as_time(value):
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, time):
        return value
    s = clean(value)
    if not s:
        return None
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(s, fmt).time()
        except ValueError:
            pass
    return None


SPANISH_DAYS = "lun|mar|mié|mie|jue|vie|sáb|sab|dom"


def as_spanish_datetime(value):
    if isinstance(value, datetime):
        return value
    s = clean(value)
    s = re.sub(rf"^(?:{SPANISH_DAYS})\s+", "", s, flags=re.I)
    for fmt in ("%d/%m/%y %H:%M", "%d/%m/%Y %H:%M"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return None


def num(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    s = clean(value).replace(" ", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def duration_hours(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    match = re.search(r"-?\d+(?:[.,]\d+)?", clean(value))
    return float(match.group().replace(",", ".")) if match else None


def load_rows(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    values = list(ws.iter_rows(values_only=True))
    headers = [clean(v) for v in values[0]]
    return headers, [dict(zip(headers, row)) for row in values[1:]]


def combine_dt(dv, tv):
    d = as_date(dv)
    t = as_time(tv)
    return datetime.combine(d, t) if d and t else None


def proposed_category(description, cause=""):
    text = normalize(description)
    rules = [
        ("Neumática, aire y compresor", ("perdida de aire", "manguera", "compresor", "sin fuerza")),
        ("Traslación", ("traslac", "no traslada", "se desplaza", "carro")),
        ("Izaje y aparejo", ("izaje", "aparejo", "no sube", "no baja")),
        ("Giro y percha", ("gira", "giro", "percha")),
        ("Herramientas, almeja y cierre", ("almeja", "alicate", "destornillador", "picador", "herramient")),
        ("Mando y control", ("mando", "control", "selectora", "remoto")),
        ("Alimentación eléctrica", ("energia", "electr", "contactor", "fusible", "cable", "falso contacto")),
        ("Lubricación", ("lubric",)),
        ("Freno", ("freno",)),
        ("Estructura o elemento mecánico", ("ruido", "fatiga", "bulon", "cabezal", "rodamiento", "engranaje")),
    ]
    for category, words in rules:
        if any(word in text for word in words):
            return category
    return "Otra / revisión humana"


def main():
    avis_headers, avis_rows = load_rows(FILES["avisos"])
    notif_headers, notif_rows = load_rows(FILES["notificaciones"])

    # Maintenance workbook has a title row before the actual headers.
    wb = openpyxl.load_workbook(FILES["mantenimiento"], read_only=True, data_only=True)
    ws = wb.worksheets[0]
    values = list(ws.iter_rows(values_only=True))
    maint_headers = [clean(v) for v in values[1]]
    maint_rows = [dict(zip(maint_headers, row)) for row in values[2:]]

    validation = []
    excluded = {"avisos": Counter(), "notificaciones": Counter(), "mantenimiento": Counter()}

    # Exact duplicate checks on every source row.
    for source, rows in (("avisos", avis_rows), ("notificaciones", notif_rows), ("mantenimiento", maint_rows)):
        sigs = Counter(tuple(clean(v) for v in row.values()) for row in rows if any(clean(v) for v in row.values()))
        for sig, count in sigs.items():
            if count > 1:
                validation.append({"tipo": "Registro duplicado exacto", "fuente": source, "cantidad": count, "detalle": "Fila repetida exactamente"})

    maint = []
    for idx, row in enumerate(maint_rows, start=3):
        tag = clean(row.get("Equipo"))
        if not tag:
            continue
        if tag not in TAG_SET:
            excluded["mantenimiento"][tag] += 1
            continue
        start = as_spanish_datetime(row.get("Comienzo"))
        finish = as_spanish_datetime(row.get("Fin"))
        hours = duration_hours(row.get("Duración") or row.get("Duraci�n"))
        if not start or not finish:
            validation.append({"tipo": "Fecha de mantenimiento inválida", "fuente": "mantenimiento", "fila": idx, "equipo": tag})
            continue
        if start.date() < START or start.date() > END:
            continue
        if hours is None or hours < 0:
            validation.append({"tipo": "Horas de mantenimiento inválidas", "fuente": "mantenimiento", "fila": idx, "equipo": tag, "valor": clean(row.get("Duración") or row.get("Duraci�n"))})
            continue
        maint.append({"fila": idx, "equipo": tag, "tarea": clean(row.get("Nombre de tarea")), "horas": hours, "inicio": start, "fin": finish, "observaciones": clean(row.get("Observaciones"))})

    notifs_by_order = defaultdict(list)
    for idx, row in enumerate(notif_rows, start=2):
        tag = clean(row.get("Equipo"))
        if tag not in TAG_SET:
            if tag:
                excluded["notificaciones"][tag] += 1
            continue
        d = as_date(row.get("Fecha de inicio real"))
        if d and START <= d <= END:
            order = clean(row.get("Orden"))
            notifs_by_order[order].append({
                "fila": idx, "equipo": tag, "contador": clean(row.get("Contador")),
                "texto": clean(row.get("Texto de notificación") or row.get("Texto de notificaci�n")),
                "trabajo_real": num(row.get("Trabajo real")),
                "notificacion_final": clean(row.get("Notificación final") or row.get("Notificaci�n final")),
            })

    events = []
    for idx, row in enumerate(avis_rows, start=2):
        tag = clean(row.get("Equipo"))
        if tag not in TAG_SET:
            if tag:
                excluded["avisos"][tag] += 1
            continue
        event_date = as_date(row.get("Fecha de aviso"))
        if not event_date or not (START <= event_date <= END):
            continue
        stopped = normalize(row.get("Parada")) == "x"
        duration = num(row.get("Duración de parada") or row.get("Duraci�n de parada"))
        valid_duration = (not stopped) or (duration is not None and duration > 0)
        if stopped and not valid_duration:
            validation.append({"tipo": "Parada sin duración válida", "fuente": "avisos", "fila": idx, "equipo": tag, "aviso": clean(row.get("Aviso")), "valor": clean(row.get("Duración de parada") or row.get("Duraci�n de parada"))})
        order = clean(row.get("Orden"))
        notif_candidates = notifs_by_order.get(order, [])
        matched = [n for n in notif_candidates if n["equipo"] == tag]
        mismatch = [n for n in notif_candidates if n["equipo"] != tag]
        if mismatch:
            validation.append({"tipo": "Orden coincide pero equipo difiere", "fuente": "vinculación", "fila": idx, "equipo": tag, "orden": order, "equipos_notificación": sorted({n["equipo"] for n in mismatch})})
        unique_texts = []
        for n in sorted(matched, key=lambda x: (x["contador"], x["fila"])):
            if n["texto"] and n["texto"] not in unique_texts:
                unique_texts.append(n["texto"])
        description = clean(row.get("Descripción") or row.get("Descripci�n"))
        cause = clean(row.get("CAUSA - Texto"))
        start_dt = combine_dt(row.get("Inicio de avería") or row.get("Inicio de aver�a"), row.get("Hora de inicio de avería") or row.get("Hora de inicio de aver�a"))
        end_dt = combine_dt(row.get("Fin de avería") or row.get("Fin de aver�a"), row.get("Hora de fin de avería") or row.get("Hora de fin de aver�a"))
        events.append({
            "fila": idx, "aviso": clean(row.get("Aviso")), "equipo": tag, "descripcion": description,
            "texto_ampliado": clean(row.get("Texto Ampl")), "parada": stopped,
            "duracion": duration if duration is not None else 0.0, "duracion_valida": valid_duration,
            "causa": cause, "orden": order, "inicio": start_dt, "fin": end_dt,
            "intervenciones": unique_texts, "notificacion_encontrada": bool(matched),
            "categoria": proposed_category(description, cause),
        })
        if order and not matched:
            validation.append({"tipo": "Orden sin notificación vinculable", "fuente": "vinculación", "fila": idx, "equipo": tag, "orden": order, "aviso": clean(row.get("Aviso"))})

    # Overlap of stops for the same bridge.
    stop_events = [e for e in events if e["parada"] and e["inicio"] and e["fin"]]
    for i, a in enumerate(stop_events):
        for b in stop_events[i + 1:]:
            if a["equipo"] == b["equipo"] and max(a["inicio"], b["inicio"]) < min(a["fin"], b["fin"]):
                validation.append({"tipo": "Paradas superpuestas", "fuente": "avisos", "equipo": a["equipo"], "avisos": [a["aviso"], b["aviso"]], "intervalos": [f"{a['inicio']:%d/%m %H:%M}-{a['fin']:%d/%m %H:%M}", f"{b['inicio']:%d/%m %H:%M}-{b['fin']:%d/%m %H:%M}"]})

    # Stop coincident with a planned maintenance interval for the same bridge.
    for e in stop_events:
        for m in maint:
            if e["equipo"] == m["equipo"] and max(e["inicio"], m["inicio"]) < min(e["fin"], m["fin"]):
                validation.append({"tipo": "Parada coincidente con mantenimiento", "fuente": "cruce avisos-mantenimiento", "equipo": e["equipo"], "aviso": e["aviso"], "tarea": m["tarea"], "intervalo_parada": f"{e['inicio']:%d/%m %H:%M}-{e['fin']:%d/%m %H:%M}", "intervalo_mantenimiento": f"{m['inicio']:%d/%m %H:%M}-{m['fin']:%d/%m %H:%M}"})

    by_tag = []
    for tag in TAGS:
        mh = sum(m["horas"] for m in maint if m["equipo"] == tag)
        tag_events = [e for e in events if e["equipo"] == tag]
        invalid = [e for e in tag_events if e["parada"] and not e["duracion_valida"]]
        sh = sum(e["duracion"] for e in tag_events if e["parada"] and e["duracion_valida"])
        available = 168.0 - mh
        availability = None if invalid or available <= 0 else (available - sh) / available * 100
        by_tag.append({
            "equipo": tag, "sala": ROOM[tag], "serie": SERIES[tag], "calendario": 168.0,
            "mantenimiento": mh, "tiempo_disponible": available, "paradas": sh,
            "detenciones": sum(1 for e in tag_events if e["parada"]), "avisos": len(tag_events),
            "disponibilidad": availability,
            "fallas": [e["descripcion"] for e in tag_events if e["parada"]],
        })

    def aggregate(name, members, calendar):
        rows = [r for r in by_tag if r["equipo"] in members]
        maintenance = sum(r["mantenimiento"] for r in rows)
        stops = sum(r["paradas"] for r in rows)
        available = calendar - maintenance
        invalid_members = [r["equipo"] for r in rows if r["disponibilidad"] is None]
        availability = None if invalid_members or available <= 0 else (available - stops) / available * 100
        return {"nombre": name, "calendario": calendar, "mantenimiento": maintenance, "tiempo_disponible": available, "paradas": stops, "disponibilidad": availability, "equipos_no_calculables": invalid_members}

    rooms = [aggregate(f"Sala {room}", [t for t in TAGS if ROOM[t] == room], 504.0) for room in range(1, 5)]
    series = [aggregate("Serie A", [t for t in TAGS if SERIES[t] == "A"], 1008.0), aggregate("Serie B", [t for t in TAGS if SERIES[t] == "B"], 1008.0)]
    global_result = aggregate("Global", TAGS, 2016.0)

    # Intra-week semantic groups are observations, not historical recurrence.
    category_groups = []
    for category in sorted({e["categoria"] for e in events if e["parada"]}):
        members = [e for e in events if e["parada"] and e["categoria"] == category]
        if len(members) >= 2:
            category_groups.append({
                "categoria": category, "cantidad": len(members), "equipos": sorted({e["equipo"] for e in members}, key=TAGS.index),
                "avisos": [e["aviso"] for e in members], "descripciones": [e["descripcion"] for e in members],
                "causas": [e["causa"] for e in members], "intervenciones": [" | ".join(e["intervenciones"]) for e in members],
            })

    output = {
        "periodo": {"semana": WEEK, "anio": YEAR, "inicio": START.isoformat(), "fin": END.isoformat(), "semanas_analizadas": 1},
        "fuentes": {k: v.name for k, v in FILES.items()},
        "conteos_fuente": {"avisos_filas": len(avis_rows), "notificaciones_filas": len(notif_rows), "mantenimiento_filas": len(maint_rows)},
        "excluidos": {source: dict(counts) for source, counts in excluded.items()},
        "puentes": by_tag, "salas": rooms, "series": series, "global": global_result,
        "avisos": [{**e, "inicio": e["inicio"].isoformat() if e["inicio"] else None, "fin": e["fin"].isoformat() if e["fin"] else None} for e in events],
        "mantenimiento": [{**m, "inicio": m["inicio"].isoformat(), "fin": m["fin"].isoformat()} for m in maint],
        "grupos_semana": category_groups, "validaciones": validation,
    }
    out_path = ROOT / "Trabajo final" / f"semana{WEEK}_datos.json"
    out_path.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "global": global_result,
        "puentes": by_tag,
        "salas": rooms,
        "series": series,
        "avisos_con_detencion": sum(1 for e in events if e["parada"]),
        "horas_detencion": sum(e["duracion"] for e in events if e["parada"] and e["duracion_valida"]),
        "validaciones": validation,
        "excluidos": output["excluidos"],
        "grupos_semana": category_groups,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--week", type=int, default=33)
    parser.add_argument("--year", type=int, default=2026)
    parser.add_argument("--start", default="2026-08-10")
    parser.add_argument("--end", default="2026-08-16")
    parser.add_argument("--avisos", default="33-Avisos MG (1).XLSX")
    parser.add_argument("--notificaciones", default="33-Notificaciones (1).XLSX")
    parser.add_argument("--mantenimiento", default="33-TAREAS EN SALA SEMANA 33 (1).xlsx")
    args = parser.parse_args()
    WEEK, YEAR = args.week, args.year
    START, END = date.fromisoformat(args.start), date.fromisoformat(args.end)
    FILES = {
        "avisos": INPUT / args.avisos,
        "notificaciones": INPUT / args.notificaciones,
        "mantenimiento": INPUT / args.mantenimiento,
    }
    main()
