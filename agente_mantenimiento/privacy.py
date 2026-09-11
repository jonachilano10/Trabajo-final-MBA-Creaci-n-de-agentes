from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterable

import openpyxl
from openpyxl.utils import get_column_letter


class ConfidentialDataError(ValueError):
    """Raised before analysis when an upload may contain personal data."""


@dataclass(frozen=True)
class ConfidentialLocation:
    filename: str
    sheet: str
    column: str
    header: str
    reason: str


SENSITIVE_HEADERS = {
    "nombre", "apellido", "nombre completo", "empleado", "operario", "operador",
    "supervisor", "legajo", "dni", "documento", "documento identidad", "correo",
    "email", "e mail", "telefono", "celular", "movil",
}
BUSINESS_ID_HEADERS = {
    "aviso", "orden", "notific", "equipo", "contador", "codigo equipo",
    "numero aviso", "numero orden", "numero notificacion",
}
COMMON_GIVEN_NAMES = {
    "juan", "maria", "jose", "ana", "carlos", "laura", "lucas", "sofia",
    "martin", "martina", "diego", "paula", "pablo", "gabriela", "miguel",
    "andrea", "fernando", "florencia", "javier", "valentina", "mariano",
    "carolina", "sebastian", "natalia", "alejandro", "daniela", "matias",
    "camila", "roberto", "silvia", "ricardo", "monica", "gustavo", "patricia",
}
EMAIL_RE = re.compile(r"(?<![\w.-])[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}(?![\w.-])")
ROLE_NAME_RE = re.compile(
    r"\b(?:operari[oa]|operador(?:a)?|supervisor(?:a)?|emplead[oa])\s*[:\-]?\s*"
    r"[A-ZÁÉÍÓÚÑ][a-záéíóúñ]+(?:\s+[A-ZÁÉÍÓÚÑ][a-záéíóúñ]+){1,3}\b"
)
PHONE_LABEL_RE = re.compile(r"\b(?:tel(?:efono)?|cel(?:ular)?|movil)\s*[:\-]?\s*\+?[\d ()-]{7,}\b", re.I)
INTERNATIONAL_PHONE_RE = re.compile(r"(?<!\w)\+\d[\d ()-]{7,}\d(?!\w)")
PERSON_NAME_RE = re.compile(r"\b([A-ZÁÉÍÓÚÑ][a-záéíóúñ]+)\s+([A-ZÁÉÍÓÚÑ][a-záéíóúñ]+)\b")
EQUIPMENT_RE = re.compile(r"P-MG(?:11|12|13|21|22|23|31|32|33|41|42|43)", re.I)


def _normalized(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def _header_kind(header: str) -> tuple[bool, bool]:
    normalized = _normalized(header)
    sensitive = any(re.search(rf"\b{re.escape(item)}\b", normalized) for item in SENSITIVE_HEADERS)
    business = any(re.search(rf"\b{re.escape(item)}", normalized) for item in BUSINESS_ID_HEADERS)
    return sensitive, business


def detect_confidential_reason(value: Any, header: str = "") -> str | None:
    if value is None or isinstance(value, (bool, date, datetime, time)):
        return None
    text = str(value).strip()
    if not text:
        return None
    sensitive_header, business_header = _header_kind(header)
    if sensitive_header:
        return "columna de datos personales"
    if EMAIL_RE.search(text):
        return "correo electrónico"
    if ROLE_NAME_RE.search(text):
        return "nombre de operario o supervisor"
    if PHONE_LABEL_RE.search(text) or INTERNATIONAL_PHONE_RE.search(text):
        return "teléfono"
    if EQUIPMENT_RE.fullmatch(text):
        return None
    digits = re.sub(r"[.\s-]", "", text)
    if digits.isdigit() and not business_header:
        if len(digits) == 5:
            return "posible número de legajo"
        if 7 <= len(digits) <= 9:
            return "posible DNI o documento"
    for match in PERSON_NAME_RE.finditer(text):
        if _normalized(match.group(1)) in COMMON_GIVEN_NAMES:
            return "posible nombre completo"
    return None


def scan_workbook(path: str | Path, *, display_name: str | None = None) -> list[ConfidentialLocation]:
    source = Path(path)
    try:
        workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    except Exception as error:
        raise ConfidentialDataError(f"{display_name or source.name}: el Excel está corrupto o no puede leerse.") from error
    locations: dict[tuple[str, str, str], ConfidentialLocation] = {}
    try:
        for worksheet in workbook.worksheets:
            headers: dict[int, str] = {}
            for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                if row_index == 1:
                    headers = {index: str(value).strip() if value is not None else "Sin encabezado"
                               for index, value in enumerate(row, start=1)}
                for column_index, value in enumerate(row, start=1):
                    raw_header = headers.get(column_index, "Sin encabezado")
                    header = "Sin encabezado" if row_index == 1 else raw_header
                    reason = detect_confidential_reason(value, header)
                    if reason:
                        column = get_column_letter(column_index)
                        key = (worksheet.title, column, reason)
                        safe_header = (
                            "Encabezado confidencial"
                            if detect_confidential_reason(raw_header, "") else raw_header
                        )
                        locations[key] = ConfidentialLocation(
                            display_name or source.name, worksheet.title, column, safe_header, reason
                        )
    finally:
        workbook.close()
    return list(locations.values())


def validate_no_confidential_data(files: Iterable[tuple[str | Path, str]]) -> None:
    findings: list[ConfidentialLocation] = []
    for path, display_name in files:
        findings.extend(scan_workbook(path, display_name=display_name))
    if findings:
        locations = "; ".join(
            f"{item.filename} - hoja {item.sheet}, columna {item.column} ({item.header}): {item.reason}"
            for item in findings
        )
        raise ConfidentialDataError(
            "Los archivos contienen información confidencial, no se ejecutará el análisis. "
            f"Eliminá los datos indicados antes de volver a subirlos: {locations}."
        )
